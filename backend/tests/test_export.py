"""Exports: JSON, CSV, XLSX, PDF generation from local data."""

from __future__ import annotations

import io
import json

from conftest import INVOICE_TEXT, make_png, process_document, upload_document


def _processed_document(client, fake_ocr):
    fake_ocr(INVOICE_TEXT)
    document = upload_document(client, make_png(), "invoice.png")
    process_document(client, document["id"])
    return document["id"]


def test_json_export(client, fake_ocr):
    document_id = _processed_document(client, fake_ocr)
    response = client.get(f"/api/export/{document_id}/json")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/json")
    assert "attachment" in response.headers["content-disposition"]

    payload = json.loads(response.content)
    assert payload["generator"] == "SecureDoc AI"
    assert payload["document"]["document_type"] == "invoice"
    names = {f["name"] for f in payload["fields"]}
    assert "invoice_number" in names
    assert payload["line_items"]
    assert payload["validation"]
    assert payload["validation_summary"]["total"] >= 1


def test_csv_export(client, fake_ocr):
    document_id = _processed_document(client, fake_ocr)
    response = client.get(f"/api/export/{document_id}/csv")
    assert response.status_code == 200
    text = response.content.decode("utf-8-sig")
    assert text.splitlines()[0].startswith("section,name,value")
    assert "field,invoice_number" in text.replace("\"", "")
    assert "INV-2026-1042" in text


def test_xlsx_export(client, fake_ocr):
    from openpyxl import load_workbook

    document_id = _processed_document(client, fake_ocr)
    response = client.get(f"/api/export/{document_id}/xlsx")
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    assert set(workbook.sheetnames) >= {"Summary", "Fields", "Line Items", "Validation"}
    fields_sheet = workbook["Fields"]
    values = [row[0].value for row in fields_sheet.iter_rows(min_row=2)]
    assert "invoice_number" in values


def test_pdf_export(client, fake_ocr):
    document_id = _processed_document(client, fake_ocr)
    response = client.get(f"/api/export/{document_id}/pdf")
    assert response.status_code == 200
    assert response.content.startswith(b"%PDF")
    assert "securedoc_invoice_" in response.headers["content-disposition"]


def test_export_unknown_format_rejected(client, fake_ocr):
    document_id = _processed_document(client, fake_ocr)
    response = client.get(f"/api/export/{document_id}/exe")
    assert response.status_code == 400
    assert response.json()["error"]["code"] == "invalid_format"


def test_export_missing_document_404(client):
    assert client.get("/api/export/nope/json").status_code == 404


def test_export_filename_is_generated_not_user_controlled(client, fake_ocr):
    document_id = _processed_document(client, fake_ocr)
    response = client.get(f"/api/export/{document_id}/json")
    disposition = response.headers["content-disposition"]
    assert "invoice.png" not in disposition
    assert document_id[:8] in disposition
