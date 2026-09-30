"""Local export builders: JSON, CSV, XLSX, PDF."""

from __future__ import annotations

import csv
import io
import json

from app.models import Classification, Document, ExtractedField, LineItem, ValidationResult


def build_payload(db, document: Document) -> dict:
    classification = (
        db.query(Classification)
        .filter(Classification.document_id == document.id)
        .order_by(Classification.id.desc())
        .first()
    )
    fields = [
        {
            "name": row.field_name,
            "value": row.value or "",
            "confidence": round(float(row.confidence or 0), 4),
            "source": row.source,
            "is_corrected": bool(row.is_corrected),
            "original_value": row.original_value,
        }
        for row in db.query(ExtractedField).filter(ExtractedField.document_id == document.id)
    ]
    fields.sort(key=lambda f: f["name"])
    line_items = [
        {
            "description": row.description,
            "quantity": row.quantity,
            "unit_price": row.unit_price,
            "tax": row.tax,
            "total": row.total,
        }
        for row in db.query(LineItem)
        .filter(LineItem.document_id == document.id)
        .order_by(LineItem.position)
    ]
    validations = []
    for row in db.query(ValidationResult).filter(ValidationResult.document_id == document.id):
        try:
            details = json.loads(row.details_json or "{}")
        except (TypeError, ValueError):
            details = {}
        validations.append(
            {
                "rule": row.rule,
                "passed": bool(row.passed),
                "message": row.message,
                "severity": details.get("severity", "error"),
            }
        )
    return {
        "generator": "SecureDoc AI",
        "note": "Generated locally. No data left this device.",
        "document": {
            "id": document.id,
            "filename": document.filename,
            "document_type": document.document_type,
            "status": document.status,
            "page_count": document.page_count,
            "overall_confidence": round(float(document.overall_confidence or 0), 4),
            "processing_time_ms": document.processing_time_ms,
            "created_at": document.created_at.isoformat() if document.created_at else None,
        },
        "classification": (
            {
                "document_type": classification.document_type,
                "confidence": round(float(classification.confidence), 4),
                "model": classification.model_name,
            }
            if classification
            else None
        ),
        "fields": fields,
        "line_items": line_items,
        "validation": validations,
        "validation_summary": {
            "passed": sum(1 for v in validations if v["passed"]),
            "total": len(validations),
        },
    }


def payload_to_json(payload: dict) -> bytes:
    return json.dumps(payload, indent=2, ensure_ascii=False).encode("utf-8")


def payload_to_csv(payload: dict, include_confidence: bool = True) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    header = ["section", "name", "value"]
    if include_confidence:
        header += ["confidence", "source", "corrected"]
    writer.writerow(header)

    def row(section: str, name: str, value, confidence=None, source="", corrected=""):
        record = [section, name, value]
        if include_confidence:
            record += [confidence, source, corrected]
        writer.writerow(record)

    doc = payload["document"]
    row("document", "id", doc["id"])
    row("document", "filename", doc["filename"])
    row("document", "type", doc["document_type"])
    row("document", "status", doc["status"])
    row("document", "overall_confidence", doc["overall_confidence"])
    for field in payload["fields"]:
        row("field", field["name"], field["value"], field["confidence"], field["source"],
            "yes" if field["is_corrected"] else "no")
    for index, item in enumerate(payload["line_items"], start=1):
        for key in ("description", "quantity", "unit_price", "tax", "total"):
            row(f"line_item_{index}", key, item[key])
    for validation in payload["validation"]:
        row("validation", validation["rule"], "PASS" if validation["passed"] else "FAIL")
    return buffer.getvalue().encode("utf-8-sig")


def payload_to_xlsx(payload: dict, include_confidence: bool = True) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"
    ws.append(["SecureDoc AI — Document Intelligence Report"])
    ws.append(["Generated locally; no data left this device"])
    ws.append([])
    doc = payload["document"]
    for key, value in (
        ("Document ID", doc["id"]),
        ("Filename", doc["filename"]),
        ("Type", doc["document_type"]),
        ("Status", doc["status"]),
        ("AI confidence", doc["overall_confidence"]),
        ("Validation", f"{payload['validation_summary']['passed']}/{payload['validation_summary']['total']} rules passed"),
    ):
        ws.append([key, value])

    ws_fields = wb.create_sheet("Fields")
    header = ["Field", "Value"]
    if include_confidence:
        header += ["Confidence", "Source", "Corrected"]
    ws_fields.append(header)
    for field in payload["fields"]:
        row = [field["name"], field["value"]]
        if include_confidence:
            row += [field["confidence"], field["source"], "yes" if field["is_corrected"] else "no"]
        ws_fields.append(row)

    ws_items = wb.create_sheet("Line Items")
    ws_items.append(["Description", "Quantity", "Unit Price", "Tax", "Total"])
    for item in payload["line_items"]:
        ws_items.append([item["description"], item["quantity"], item["unit_price"],
                         item["tax"], item["total"]])

    ws_valid = wb.create_sheet("Validation")
    ws_valid.append(["Rule", "Result", "Message"])
    for validation in payload["validation"]:
        ws_valid.append([validation["rule"], "PASS" if validation["passed"] else "FAIL",
                         validation["message"]])

    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


def payload_to_pdf(payload: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    stream = io.BytesIO()
    doc = SimpleDocTemplate(stream, pagesize=A4, topMargin=18 * mm, bottomMargin=18 * mm)
    styles = getSampleStyleSheet()
    title = ParagraphStyle("Title2", parent=styles["Title"], fontSize=20, spaceAfter=4,
                           textColor=colors.HexColor("#0f172a"))
    subtitle = ParagraphStyle("Sub", parent=styles["Normal"], fontSize=9,
                              textColor=colors.HexColor("#64748b"), spaceAfter=14)
    heading = ParagraphStyle("Head", parent=styles["Heading3"], fontSize=12, spaceBefore=12,
                             spaceAfter=6, textColor=colors.HexColor("#0f172a"))
    body = ParagraphStyle("Body2", parent=styles["Normal"], fontSize=9.5)

    story = [
        Paragraph("SECUREDOC AI", title),
        Paragraph("Document Intelligence Report — generated locally on this device.", subtitle),
    ]

    meta = payload["document"]
    summary_rows = [
        ["Document", meta["filename"]],
        ["Type", meta["document_type"]],
        ["Status", meta["status"]],
        ["AI confidence", f"{round(meta['overall_confidence'] * 100)}%"],
        ["Validation",
         f"{payload['validation_summary']['passed']}/{payload['validation_summary']['total']} rules passed"],
    ]
    classification = payload.get("classification")
    if classification:
        summary_rows.append(["Classifier", f"{classification['model']} "
                                          f"({round(classification['confidence'] * 100)}%)"])
    story.append(_table(summary_rows, body))
    story.append(Spacer(1, 8))

    story.append(Paragraph("Extracted Fields", heading))
    rows = [["Field", "Value", "AI confidence"]]
    for field in payload["fields"]:
        rows.append([field["name"], field["value"],
                     f"{round(field['confidence'] * 100)}%" + (" ✓ corrected" if field["is_corrected"] else "")])
    story.append(_table(rows, body))

    if payload["line_items"]:
        story.append(Paragraph("Line Items", heading))
        rows = [["Description", "Qty", "Rate", "Tax", "Total"]]
        for item in payload["line_items"]:
            rows.append([item["description"], item["quantity"], item["unit_price"],
                         item["tax"], item["total"]])
        story.append(_table(rows, body))

    story.append(Paragraph("Validation", heading))
    rows = [["Rule", "Result", "Message"]]
    for validation in payload["validation"]:
        rows.append([validation["rule"],
                     "PASS" if validation["passed"] else "FAIL",
                     validation["message"]])
    story.append(_table(rows, body))

    story.append(Spacer(1, 12))
    story.append(Paragraph(
        "Confidence reflects AI certainty, not a guarantee of correctness. "
        "This report was produced fully offline by SecureDoc AI.", subtitle,
    ))
    doc.build(story)
    return stream.getvalue()


def _table(rows, body_style):
    from reportlab.lib import colors
    from reportlab.platypus import Table, TableStyle

    rendered = Table(rows, hAlign="LEFT")
    rendered.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#334155")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2e8f0")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    return rendered
